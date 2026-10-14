"""Daily digest assembly and Excel export."""

from __future__ import annotations

import html
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from .config import Config, now_ist
from .storage import Bid, latest_scores, last_successful_digest, was_notified


def fmt_inr(amount: float | None) -> str:
    """Indian digit grouping: 1234567 -> ₹12,34,567."""
    if amount is None:
        return "-"
    whole = f"{int(round(amount))}"
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join(groups + [tail])
    return f"₹{whole}"


@dataclass
class DigestEntry:
    bid_id: int
    bid_number: str
    title: str
    organisation: str
    closing: datetime | None
    days_left: float | None
    score: float
    eligibility: str
    reasons: list[str]
    emd: str
    link: str
    change: str | None = None
    feedback: str | None = None

    @property
    def closing_str(self) -> str:
        return self.closing.strftime("%d %b %Y %H:%M") if self.closing else "-"


@dataclass
class Digest:
    date: datetime
    new: list[DigestEntry] = field(default_factory=list)
    closing: list[DigestEntry] = field(default_factory=list)
    changed: list[DigestEntry] = field(default_factory=list)
    stats: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    dashboard_url: str = ""

    @property
    def is_empty(self) -> bool:
        return not (self.new or self.closing or self.changed)

    @property
    def subject(self) -> str:
        return (
            f"GeM digest {self.date:%d %b}: {len(self.new)} new, "
            f"{len(self.closing)} closing soon, {len(self.changed)} changed"
        )

    def all_entries(self) -> list[DigestEntry]:
        return self.new + self.closing + self.changed


def _emd_text(bid: Bid, cfg: Config) -> str:
    doc = bid.document
    if doc is None or not doc.parsed_ok:
        return "-"
    if not doc.emd_amount:
        return "No EMD"
    text = fmt_inr(doc.emd_amount)
    if cfg.company.mse_benefits or cfg.company.startup_valid_on(bid.end_date):
        text += " (exemption likely)"
    return text


def _entry(bid: Bid, score, cfg: Config, now: datetime) -> DigestEntry:
    days = (bid.end_date - now).total_seconds() / 86400 if bid.end_date else None
    return DigestEntry(
        bid_id=bid.id,
        bid_number=bid.bid_number,
        title=bid.title,
        organisation=" / ".join(x for x in [bid.organisation or bid.department, bid.ministry] if x) or "-",
        closing=bid.end_date,
        days_left=days,
        score=score.total if score else 0.0,
        eligibility=score.eligibility_status if score else "-",
        reasons=list(score.eligibility_reasons or []) if score else [],
        emd=_emd_text(bid, cfg),
        link=bid.doc_url,
        change=bid.change_summary,
        feedback=bid.feedback,
    )


def build_digest(session: Session, cfg: Config, run_id: int | None = None, now: datetime | None = None) -> Digest:
    now = now or now_ist()
    n = cfg.notifications
    digest = Digest(date=now, dashboard_url=n.dashboard_url)
    since = last_successful_digest(session) or (now - timedelta(days=1))

    candidates = list(
        session.scalars(
            select(Bid)
            .options(selectinload(Bid.document))
            .where(Bid.shortlisted_at.is_not(None), Bid.end_date > now)
        )
    )
    scores = latest_scores(session, [b.id for b in candidates])

    for bid in candidates:
        score = scores.get(bid.id)
        entry = _entry(bid, score, cfg, now)
        acted = bid.feedback in ("applied", "not_relevant")
        is_new = bool(score and score.shortlisted and not acted and not was_notified(session, bid.id, "new"))
        if is_new:
            digest.new.append(entry)
        elif (
            entry.days_left is not None
            and entry.days_left <= n.closing_soon_days
            and not acted
            and bid.gem_status == "active"
        ):
            digest.closing.append(entry)
        if not is_new and bid.changed_at and bid.changed_at > since and bid.feedback != "not_relevant":
            digest.changed.append(entry)

    digest.new.sort(key=lambda e: e.score, reverse=True)
    digest.closing.sort(key=lambda e: e.closing or now)
    digest.changed.sort(key=lambda e: e.score, reverse=True)
    cap = n.max_items_per_section
    digest.new, digest.closing, digest.changed = digest.new[:cap], digest.closing[:cap], digest.changed[:cap]
    return digest


# -- rendering --------------------------------------------------------------------------


def _days(e: DigestEntry) -> str:
    return "-" if e.days_left is None else f"{e.days_left:.0f}d"


def _eligibility(e: DigestEntry) -> str:
    return f"{e.eligibility} ({'; '.join(e.reasons)})" if e.reasons else e.eligibility


SECTIONS = [
    ("new", "New relevant bids"),
    ("closing", "Closing soon - not yet acted on"),
    ("changed", "Changed since last digest"),
]


def render_text(d: Digest) -> str:
    lines = [d.subject, ""]
    for key, heading in SECTIONS:
        entries: list[DigestEntry] = getattr(d, key)
        lines.append(f"== {heading} ({len(entries)}) ==")
        if not entries:
            lines.append("  none")
        for i, e in enumerate(entries, 1):
            lines.append(f"{i}. [{e.score:.0f}] {e.bid_number} - {e.title}")
            lines.append(f"   {e.organisation}")
            lines.append(f"   Closes {e.closing_str} ({_days(e)}) | {_eligibility(e)} | EMD {e.emd}")
            if key == "changed" and e.change:
                lines.append(f"   Change: {e.change}")
            lines.append(f"   {e.link}")
        lines.append("")
    lines += [f"Warning: {w}" for w in d.warnings]
    if d.dashboard_url:
        lines.append(f"Dashboard: {d.dashboard_url}")
    return "\n".join(lines)


def render_html(d: Digest) -> str:
    esc = html.escape
    parts = [
        "<html><body style='font-family:Segoe UI,Arial,sans-serif;font-size:14px;color:#222'>",
        f"<h2 style='margin-bottom:4px'>{esc(d.subject)}</h2>",
    ]
    for w in d.warnings:
        parts.append(f"<p style='color:#b45309'><b>Warning:</b> {esc(w)}</p>")
    for key, heading in SECTIONS:
        entries: list[DigestEntry] = getattr(d, key)
        parts.append(f"<h3 style='margin:18px 0 6px'>{esc(heading)} ({len(entries)})</h3>")
        if not entries:
            parts.append("<p style='color:#666'>None.</p>")
            continue
        parts.append(
            "<table cellpadding='6' style='border-collapse:collapse;width:100%'>"
            "<tr style='background:#f1f5f9;text-align:left'><th>Score</th><th>Bid</th><th>Organisation</th>"
            "<th>Closes</th><th>Eligibility</th><th>EMD</th></tr>"
        )
        for e in entries:
            change = f"<br><i style='color:#b45309'>{esc(e.change)}</i>" if key == "changed" and e.change else ""
            parts.append(
                "<tr style='border-top:1px solid #e2e8f0;vertical-align:top'>"
                f"<td><b>{e.score:.0f}</b></td>"
                f"<td><a href='{esc(e.link)}'>{esc(e.bid_number)}</a><br>{esc(e.title[:160])}{change}</td>"
                f"<td>{esc(e.organisation)}</td>"
                f"<td>{esc(e.closing_str)}<br><span style='color:#666'>{_days(e)}</span></td>"
                f"<td>{esc(_eligibility(e))}</td><td>{esc(e.emd)}</td></tr>"
            )
        parts.append("</table>")
    if d.dashboard_url:
        parts.append(f"<p><a href='{esc(d.dashboard_url)}'>Open dashboard</a></p>")
    parts.append("</body></html>")
    return "".join(parts)


def render_telegram(d: Digest) -> list[str]:
    """Telegram HTML-mode messages, each under the 4096 character limit."""
    esc = html.escape
    blocks = [f"<b>{esc(d.subject)}</b>"] + [f"⚠️ {esc(w)}" for w in d.warnings]
    for key, heading in SECTIONS:
        entries: list[DigestEntry] = getattr(d, key)
        blocks.append(f"\n<b>{esc(heading)} ({len(entries)})</b>")
        for e in entries:
            text = (
                f"<b>{e.score:.0f}</b> · <a href=\"{esc(e.link)}\">{esc(e.bid_number)}</a>\n"
                f"{esc(e.title[:150])}\n{esc(e.organisation[:80])}\n"
                f"Closes {esc(e.closing_str)} ({_days(e)}) · {esc(_eligibility(e)[:150])} · EMD {esc(e.emd)}"
            )
            if key == "changed" and e.change:
                text += f"\n<i>{esc(e.change)}</i>"
            blocks.append(text)
    messages, current = [], ""
    for block in blocks:
        if len(current) + len(block) + 2 > 3900:
            messages.append(current)
            current = ""
        current += block + "\n\n"
    if current.strip():
        messages.append(current)
    return messages


def render_markdown(d: Digest, limit: int = 3500) -> str:
    """Compact plain/markdown text for Slack and WhatsApp."""
    lines = [f"*{d.subject}*"] + [f"Warning: {w}" for w in d.warnings]
    for key, heading in SECTIONS:
        entries: list[DigestEntry] = getattr(d, key)
        lines.append(f"\n*{heading} ({len(entries)})*")
        for e in entries:
            lines.append(f"• {e.score:.0f} | {e.bid_number} | {e.title[:80]} | closes {e.closing_str} | {e.eligibility}")
            lines.append(f"  {e.link}")
    text = "\n".join(lines)
    return text if len(text) <= limit else text[: limit - 40] + "\n… (see email/dashboard for the rest)"


# -- Excel ------------------------------------------------------------------------------

COLUMNS = [
    ("Score", 8), ("Bid number", 22), ("Title", 60), ("Organisation", 40), ("Closing", 18),
    ("Days left", 9), ("Eligibility", 30), ("Notes", 50), ("EMD", 22), ("Change", 40), ("Link", 45),
]  # fmt: skip


def export_excel(d: Digest, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for key, heading in SECTIONS:
        ws = wb.create_sheet(heading[:31])
        for col, (name, width) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            ws.column_dimensions[get_column_letter(col)].width = width
        for row, e in enumerate(getattr(d, key), 2):
            values = [
                e.score, e.bid_number, e.title, e.organisation, e.closing,
                None if e.days_left is None else round(e.days_left, 1),
                e.eligibility, "; ".join(e.reasons), e.emd, e.change or "", e.link,
            ]  # fmt: skip
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=col in (3, 4, 8, 10))
            ws.cell(row=row, column=5).number_format = "dd-mmm-yyyy hh:mm"
            link = ws.cell(row=row, column=11)
            link.hyperlink = e.link
            link.font = Font(color="0563C1", underline="single")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, ws.max_row)}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
