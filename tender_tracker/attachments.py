"""Fetch documents linked from a bid (SOW, BoQ, specifications) and extract their text.

Only hosts whose robots.txt allows it are fetched. In practice
``mkp.gem.gov.in`` (BoQ/specification uploads) and ``fulfilment.gem.gov.in``
(service/SOW documents) are allowed, while buyer uploads under
``bidplus.gem.gov.in/resources/`` are disallowed. Those stay as links that
the user opens in their own browser.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import tempfile
from html.parser import HTMLParser
from pathlib import Path

from sqlalchemy.orm import Session

from .config import now_ist
from .fetcher import FetchError, GemClient
from .pdf_parser import clean_text, read_pdf
from .storage import Bid, BidAttachment

log = logging.getLogger(__name__)

# Kinds worth reading for scope of work / eligibility; GTC and policy documents are generic.
TEXT_KINDS = ("sow", "spec", "boq", "eligibility", "buyer_doc", "sla", "other")
MAX_TEXT_CHARS = 60_000


def _pdf_text(content: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as fh:
        fh.write(content)
        path = Path(fh.name)
    try:
        doc = read_pdf(path, max_pages=40)
        return clean_text(doc.text)
    finally:
        path.unlink(missing_ok=True)


def _csv_text(content: bytes) -> str:
    raw = content.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(raw)))[:500]
    return "\n".join(" | ".join(cell.strip() for cell in row if cell.strip()) for row in rows if any(row))


def _xlsx_text(content: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets[:5]:
        lines.append(f"## {ws.title}")
        for row in ws.iter_rows(max_row=500, values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


class _HtmlText(HTMLParser):
    SKIP = {"script", "style", "head", "noscript", "svg"}
    BLOCK = {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "h5", "table", "section"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _html_text(content: bytes) -> str:
    parser = _HtmlText()
    parser.feed(content.decode("utf-8", errors="replace"))
    lines = [re.sub(r"\s+", " ", ln).strip(" |") for ln in "".join(parser.parts).splitlines()]
    lines = [ln for ln in lines if ln]
    # GeM pages start with the site's navigation menu; the document body begins at its title.
    start = next(
        (i for i, ln in enumerate(lines) if re.match(r"(Special Terms|Service Level Agreement|Scope of Work)", ln, re.I)),
        0,
    )
    return "\n".join(lines[start:])


def extract_document_text(content: bytes, content_type: str, url: str) -> str | None:
    """Text of a PDF/CSV/XLSX/HTML document, or None if the format isn't supported."""
    low = url.lower().split("?")[0]
    ctype = (content_type or "").lower()
    if content.startswith(b"%PDF"):
        return _pdf_text(content)
    if "html" in ctype:
        return _html_text(content)
    if content.startswith(b"PK") and ("sheet" in ctype or low.endswith((".xlsx", ".xlsm"))):
        return _xlsx_text(content)
    if low.endswith(".csv") or "csv" in ctype:
        return _csv_text(content)
    return None


def fetch_attachment(client: GemClient, att: BidAttachment) -> None:
    att.fetched_at = now_ist()
    if not client.robots_allowed(att.url):
        att.status, att.error = "blocked_by_robots", "The host's robots.txt disallows automated access; open the link."
        return
    try:
        result = client.download(att.url)
    except FetchError as exc:
        att.status, att.error = "error", str(exc)
        return
    if result is None:
        att.status, att.error = "unavailable", "Not available (closed bid, removed, or larger than 15 MB)."
        return
    content, ctype = result
    att.content_type = ctype[:128] or None
    try:
        text = extract_document_text(content, ctype, att.url)
    except Exception as exc:  # corrupt files
        log.warning("Could not extract %s: %s", att.url, exc)
        att.status, att.error = "error", f"Could not read file: {exc}"
        return
    if text is None:
        att.status, att.error = "unsupported", f"Unsupported file type ({ctype or 'unknown'})."
        return
    if len(text.strip()) < 20:
        att.status, att.error = "unsupported", "No extractable text (scanned document?)."
        return
    att.text = text[:MAX_TEXT_CHARS]
    att.status, att.error = "fetched", None


def fetch_bid_attachments(
    session: Session, client: GemClient, bid: Bid, limit: int = 10, retry: bool = False
) -> int:
    """Fetch text for a bid's useful attachments. Returns the number of downloads attempted
    (links blocked by robots.txt cost no request and don't count)."""
    todo = [
        a
        for a in bid.attachments
        if a.kind in TEXT_KINDS
        and (a.status == "not_fetched" or (retry and a.status in ("error", "unavailable", "unsupported")))
    ]
    downloads = 0
    for att in todo:
        if downloads >= limit:
            break
        fetch_attachment(client, att)
        session.commit()
        log.info("Attachment %s (%s): %s", att.label, bid.bid_number, att.status)
        if att.status != "blocked_by_robots":
            downloads += 1
    return downloads
