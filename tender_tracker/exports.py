"""Export a proposal workspace as Word (proposal), PowerPoint (presentation) and Excel (matrices)."""

from __future__ import annotations

import io
import re
from typing import Any

from .config import CompanyProfile

STATUS_LABEL = {
    "complied": "Complied",
    "partially_complied": "Partially complied",
    "not_complied": "Not complied",
    "to_confirm": "To confirm",
}
BRAND = (0x1E, 0x3A, 0x8A)


def _text(value: Any) -> str:
    return "" if value is None else str(value)


# -- Word ------------------------------------------------------------------------------------


def _markdown_into(doc: Any, text: str) -> None:
    """Minimal markdown: #/## headings, - bullets, **bold** spans, paragraphs."""
    for line in (text or "").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")) + 1, 4)
            doc.add_heading(stripped.lstrip("#").strip(), level=level)
            continue
        bullet = re.match(r"^([-*]|\d+[.)])\s+(.*)", stripped)
        para = doc.add_paragraph(style="List Bullet" if bullet else None)
        body = bullet.group(2) if bullet else stripped
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", body)):
            run = para.add_run(part)
            run.bold = i % 2 == 1


def _table(doc: Any, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None) -> None:
    from docx.shared import Cm, Pt

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = _text(value)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph()


def to_docx(result: dict[str, Any], title: str, company: CompanyProfile) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    summary = result.get("summary", {})
    proposal = result.get("proposal") or {}
    heading = doc.add_heading("Technical Proposal", level=0)
    heading.runs[0].font.color.rgb = RGBColor(*BRAND)
    doc.add_paragraph(title).runs[0].bold = True
    meta = [
        ("Submitted by", company.name),
        ("Tender reference", summary.get("bid_number") or (result.get("key_facts") or {}).get("tender_reference")),
        ("Buyer", summary.get("organisation") or (result.get("key_facts") or {}).get("buyer")),
        ("GSTIN", company.gstin), ("PAN", company.pan), ("Udyam", company.udyam_number), ("DPIIT", company.dpiit_number),
    ]  # fmt: skip
    for label, value in meta:
        if value:
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(_text(value))
    if not proposal:
        doc.add_paragraph(
            "AI drafting was not run for this workspace, so the narrative sections are empty. "
            "The annexures below are generated from the tender documents."
        )
    doc.add_page_break()

    sections = [
        ("1. Executive summary", proposal.get("executive_summary")),
        ("2. Understanding of the requirement", proposal.get("understanding")),
    ]
    for head, body in sections:
        if body:
            doc.add_heading(head, level=1)
            _markdown_into(doc, body)

    scope = result.get("scope")
    if scope:
        doc.add_heading("3. Scope of work", level=1)
        _markdown_into(doc, scope.get("summary", ""))
        for label, key in (("Objectives", "objectives"), ("Deliverables", "deliverables")):
            if scope.get(key):
                doc.add_heading(label, level=2)
                for item in scope[key]:
                    doc.add_paragraph(item, style="List Bullet")

    arch = proposal.get("solution_architecture")
    if arch:
        doc.add_heading("4. Solution architecture", level=1)
        _markdown_into(doc, arch.get("overview", ""))
        if arch.get("components"):
            _table(doc, ["Component", "Description", "Technology"],
                   [[c["name"], c["description"], c["technology"]] for c in arch["components"]], [4, 9, 4])  # fmt: skip
    if proposal.get("technical_approach"):
        doc.add_heading("5. Technical approach", level=1)
        _markdown_into(doc, proposal["technical_approach"])
    if proposal.get("methodology"):
        doc.add_heading("6. Project methodology and plan", level=1)
        _table(doc, ["Phase", "Duration", "Activities", "Deliverables"],
               [[m["phase"], m["duration"], "\n".join(m["activities"]), "\n".join(m["deliverables"])] for m in proposal["methodology"]],
               [3.5, 2.5, 6.5, 5])  # fmt: skip
    if proposal.get("team"):
        doc.add_heading("7. Team and responsibilities", level=1)
        _table(doc, ["Role", "Count", "Responsibilities"], [[t["role"], t["count"], t["responsibilities"]] for t in proposal["team"]], [4, 1.5, 11])
    if result.get("risks"):
        doc.add_heading("8. Risks and mitigation", level=1)
        _table(doc, ["Risk", "Likelihood", "Impact", "Mitigation"],
               [[r["risk"], r["likelihood"], r["impact"], r["mitigation"]] for r in result["risks"]], [6, 2, 2, 7])  # fmt: skip
    if proposal.get("why_us"):
        doc.add_heading("9. Why us", level=1)
        for item in proposal["why_us"]:
            doc.add_paragraph(item, style="List Bullet")
    if proposal.get("assumptions"):
        doc.add_heading("10. Assumptions", level=1)
        for item in proposal["assumptions"]:
            doc.add_paragraph(item, style="List Bullet")

    if result.get("compliance"):
        doc.add_page_break()
        doc.add_heading("Annexure A - Compliance matrix", level=1)
        _table(doc, ["#", "Requirement", "Status", "Our response", "Evidence"],
               [[r["id"], r["requirement"], STATUS_LABEL.get(r["status"], r["status"]), r["response"], ", ".join(r["evidence"])]
                for r in result["compliance"]], [1.2, 5.5, 2.3, 6, 2.5])  # fmt: skip
    boq = (result.get("boq") or {}).get("items") or []
    if boq:
        doc.add_heading("Annexure B - Bill of quantities", level=1)
        _table(doc, ["#", "Item", "Description", "Qty", "Unit"],
               [[b["item_no"], b["title"], b.get("description", ""), b.get("quantity_text") or _text(b.get("quantity")), b.get("unit", "")] for b in boq],
               [1.2, 5, 7, 1.8, 2])  # fmt: skip
    required = (result.get("documents") or {}).get("required") or []
    if required:
        doc.add_heading("Annexure C - Documents to submit", level=1)
        status = {"available": "In library", "expired": "Expired", "missing": "Missing", "prepare": "Prepare", "review": "Check ATC"}
        _table(doc, ["Required document", "Status", "Our document"],
               [[d["requirement"], status.get(d["status"], d["status"]), ", ".join(x["title"] for x in d["documents"])] for d in required],
               [8, 3, 6.5])  # fmt: skip
    checklist = result.get("checklist") or []
    if checklist:
        doc.add_heading("Annexure D - Submission checklist", level=1)
        _table(doc, ["", "Item", "Detail"],
               [["✓" if c["status"] == "done" else ("!" if c["status"] == "missing" else "☐"), c["item"], c["detail"]] for c in checklist],
               [1, 7, 9.5])  # fmt: skip

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# -- PowerPoint --------------------------------------------------------------------------------


def to_pptx(result: dict[str, Any], title: str, company: CompanyProfile) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blue = RGBColor(*BRAND)

    def title_bar(slide: Any, text: str) -> None:
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = blue
        bar.line.fill.background()
        tf = bar.text_frame
        tf.text = text
        tf.paragraphs[0].runs[0].font.size = Pt(28)
        tf.paragraphs[0].runs[0].font.bold = True
        tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tf.margin_left = Inches(0.5)

    def bullets_slide(heading: str, bullets: list[str], notes: str = "") -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title_bar(slide, heading)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.6))
        tf = box.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = f"•  {b}"
            para.font.size = Pt(20 if len(bullets) <= 5 else 17)
            para.space_after = Pt(10)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    bg = cover.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = blue
    bg.line.fill.background()
    box = cover.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = "Technical Presentation"
    tf.paragraphs[0].font.size = Pt(40)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for text, size in ((title[:180], 20), (company.name, 22)):
        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0xDB, 0xE6, 0xFF)

    slides = (result.get("proposal") or {}).get("slides") or []
    if slides:
        for s in slides:
            bullets_slide(s["title"], s["bullets"], s.get("notes", ""))
    else:
        scope = result.get("scope") or {}
        if scope:
            bullets_slide("Scope of work", scope.get("deliverables") or [scope.get("summary", "")])
        boq = (result.get("boq") or {}).get("items") or []
        if boq:
            bullets_slide("Bill of quantities", [f"{b['title']} - {b.get('quantity_text') or ''} {b.get('unit') or ''}".strip(" -") for b in boq[:8]])
        bullets_slide("Next steps", ["Add ANTHROPIC_API_KEY to generate the full presentation draft"])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# -- Excel -------------------------------------------------------------------------------------


def to_xlsx(result: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    fill = PatternFill("solid", fgColor="1E3A8A")

    def sheet(name: str, headers: list[tuple[str, int]], rows: list[list[Any]]) -> None:
        ws = wb.create_sheet(name[:31])
        for col, (h, width) in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill
            ws.column_dimensions[get_column_letter(col)].width = width
        for r, row in enumerate(rows, 2):
            for col, value in enumerate(row, 1):
                ws.cell(row=r, column=col, value=value).alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"

    sheet("Compliance matrix", [("#", 6), ("Requirement", 60), ("Category", 14), ("Mandatory", 10), ("Status", 18),
                                ("Our response", 60), ("Evidence", 30), ("Remarks", 40), ("Source", 25)],
          [[r["id"], r["requirement"], r["category"], "Yes" if r["mandatory"] else "No", STATUS_LABEL.get(r["status"], r["status"]),
            r["response"], ", ".join(r["evidence"]), r["remarks"], r["source"]] for r in result.get("compliance") or []])  # fmt: skip
    sheet("Requirements", [("#", 6), ("Requirement", 80), ("Category", 14), ("Mandatory", 10), ("Source", 30)],
          [[r["id"], r["requirement"], r["category"], "Yes" if r["mandatory"] else "No", r["source"]] for r in result.get("requirements") or []])  # fmt: skip
    sheet("BoQ", [("#", 6), ("Item", 40), ("Description", 70), ("Quantity", 10), ("Unit", 10), ("Delivery", 14), ("Source", 25)],
          [[b["item_no"], b["title"], b.get("description", ""), b.get("quantity") if b.get("quantity") is not None else b.get("quantity_text"),
            b.get("unit", ""), b.get("delivery", ""), b.get("source", "")] for b in (result.get("boq") or {}).get("items") or []])  # fmt: skip
    sheet("Documents", [("Required document", 50), ("Category", 30), ("Status", 12), ("In library", 50)],
          [[d["requirement"], d["category_label"], d["status"], ", ".join(x["title"] for x in d["documents"])]
           for d in (result.get("documents") or {}).get("required") or []])  # fmt: skip
    sheet("Checklist", [("Item", 55), ("Category", 14), ("Status", 10), ("Detail", 70)],
          [[c["item"], c["category"], c["status"], c["detail"]] for c in result.get("checklist") or []])
    sheet("Risks", [("Risk", 60), ("Category", 14), ("Likelihood", 11), ("Impact", 10), ("Mitigation", 70)],
          [[r["risk"], r["category"], r["likelihood"], r["impact"], r["mitigation"]] for r in result.get("risks") or []])  # fmt: skip
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
