"""Proposal workspace pipeline with a fake AI drafter (no network, no API key)."""

import io
from datetime import datetime, timedelta

import pytest

from tender_tracker.ai import AIError, Usage
from tender_tracker.boq import boq_from_listing, find_boq
from tender_tracker.config import config_from_dict
from tender_tracker.exports import to_docx, to_pptx, to_xlsx
from tender_tracker.fetcher import parse_listing_doc
from tender_tracker.library import add_document, categorise_requirement, match_documents
from tender_tracker.pdf_parser import BidDocumentInfo
from tender_tracker.proposal import ANALYSE_SCHEMA, COMPLIANCE_SCHEMA, PROPOSAL_SCHEMA, run_analysis
from tender_tracker.storage import BidAttachment, TenderAnalysis, init_db, make_engine, save_document, session_factory, upsert_listing

from .test_fetcher_storage import DOC

ANALYSE = {
    "requirements": [
        {"id": "R1", "category": "eligibility", "requirement": "Average turnover of Rs 2.3 lakh", "mandatory": True, "source": "Bid document"},
        {"id": "R2", "category": "technical", "requirement": "Role based access control", "mandatory": True, "source": "Specification"},
    ],
    "eligibility_criteria": [{"criterion": "Turnover", "requirement": "Rs 2.3 lakh", "source": "Bid document"}],
    "mandatory_documents": [{"name": "Undertaking of non-blacklisting", "purpose": "Declaration", "source": "ATC"}],
    "boq": [],
    "scope": {"summary": "Recruitment software", "objectives": ["Digitise"], "deliverables": ["Software"], "timeline": "90 days", "location": "Bengaluru",
              "facts": [{"statement": "11 resources", "basis": "stated", "document": "GeM bid document", "page": "5", "quote": "Number of Resources : 11", "note": ""}]},
    "key_facts": {k: "" for k in ANALYSE_SCHEMA["properties"]["key_facts"]["properties"]},
    "risks": [{"risk": "Past performance gap", "category": "eligibility", "likelihood": "high", "impact": "high", "mitigation": "Seek relaxation"}],
}
COMPLIANCE = {"rows": [
    {"id": "R1", "status": "complied", "response": "Our turnover exceeds this.", "evidence": ["CA certificate FY24"], "remarks": ""},
    {"id": "R2", "status": "to_confirm", "response": "RBAC supported.", "evidence": [], "remarks": "Confirm module"},
]}
PROPOSAL = {
    "executive_summary": "We propose...", "understanding": "## Context\nThe Army needs...",
    "solution_architecture": {"overview": "Three tiers", "components": [{"name": "Portal", "description": "Web UI", "technology": "React"}]},
    "technical_approach": "## Approach\n- **Agile** delivery", "methodology": [{"phase": "Build", "duration": "6 weeks", "activities": ["Code"], "deliverables": ["App"]}],
    "team": [{"role": "PM", "count": 1, "responsibilities": "Delivery"}], "why_us": ["Startup agility"], "assumptions": ["Hosting by buyer"],
    "slides": [{"title": "Our solution", "bullets": ["Point 1", "Point 2"], "notes": "Say this"}],
}  # fmt: skip


class FakeDrafter:
    def __init__(self, fail_on=None):
        self.usage = Usage()
        self.calls = []
        self.fail_on = fail_on

    def structured(self, *, system, context, instruction, schema, max_tokens=64000):
        name = {id(ANALYSE_SCHEMA): "analyse", id(COMPLIANCE_SCHEMA): "compliance", id(PROPOSAL_SCHEMA): "proposal"}[id(schema)]
        self.calls.append((name, context))
        if name == self.fail_on:
            raise AIError("simulated failure")
        return {"analyse": ANALYSE, "compliance": COMPLIANCE, "proposal": PROPOSAL}[name]


@pytest.fixture
def env(tmp_path):
    cfg = config_from_dict({
        "company": {"name": "Acme Softech", "is_startup": True, "is_mse": True, "msme_category": "small", "turnover_lakh": 40},
        "storage": {"database_url": f"sqlite:///{(tmp_path / 't.db').as_posix()}", "data_dir": tmp_path.as_posix()},
    })  # fmt: skip
    engine = make_engine(cfg.storage.database_url)
    init_db(engine)
    Session = session_factory(engine)
    with Session() as s:
        end = (datetime.now() + timedelta(days=10)).strftime("%Y-%m-%dT%H:%M:%SZ")
        bid, _, _ = upsert_listing(s, parse_listing_doc({**DOC, "final_end_date_sort": [end]}))
        save_document(s, bid, BidDocumentInfo(
            bid_number=bid.bid_number, min_turnover_lakh=2.3, experience_years=3, emd_amount=50000,
            startup_relaxation=True, startup_relaxation_scope="Complete",
            documents_required="Experience Criteria,Bidder Turnover,Additional Doc 1 (Requested in ATC)",
        ), None)  # fmt: skip
        bid.attachments.append(BidAttachment(url="https://mkp.gem.gov.in/boq.csv", label="BoQ line items (CSV)", kind="boq", status="fetched",
            text="x", table=[["Item Number", "Item Title", "Item Description", "Item Quantity", "Unit of Measure"],
                             ["1", "Recruitment portal", "Online application module", "1", "Nos"]]))  # fmt: skip
        s.commit()
        add_document(s, cfg, kind="capability", category="turnover", title="CA certificate FY24", filename="ca.txt",
                     content=b"Turnover certificate for FY 2023-24", valid_until="2030-03-31")  # fmt: skip
        add_document(s, cfg, kind="proposal", category="proposal", title="HRMS for State PSU", filename="p.txt",
                     content=b"Recruitment Management Software proposal with role based access", outcome="won")  # fmt: skip
        a = TenderAnalysis(bid_id=bid.id, source="gem", title=bid.title)
        s.add(a)
        s.commit()
        return cfg, Session, a.id


def _run(env, drafter):
    cfg, Session, aid = env
    run_analysis(aid, cfg, Session, drafter=drafter)
    with Session() as s:
        return s.get(TenderAnalysis, aid)


def test_workspace_without_ai(env):
    a = _run(env, None)
    r = a.result
    assert a.status == "done" and not a.ai_used
    assert r["eligibility"]["status"] == "Eligible via startup/MSE exemption"
    assert r["boq"]["items"][0]["title"] == "Recruitment portal" and r["boq"]["items"][0]["quantity"] == 1
    docs = {d["requirement"]: d for d in r["documents"]["required"]}
    assert docs["Bidder Turnover"]["status"] == "available"
    assert docs["Experience Criteria"]["status"] == "missing"
    assert docs["Additional Doc 1 (Requested in ATC)"]["status"] == "review"
    assert r["stages"]["analyse"]["status"] == "skipped"
    items = {c["item"] for c in r["checklist"]}
    assert "EMD (bid security)" in items and "Price bid / BoQ rates filled and checked" in items
    assert isinstance(r["summary"]["bid_end"], str) and "T" in r["summary"]["bid_end"]  # ISO for the UI


def test_workspace_with_ai(env):
    drafter = FakeDrafter()
    a = _run(env, drafter)
    r = a.result
    assert a.ai_used and [c[0] for c in drafter.calls] == ["analyse", "compliance", "proposal"]
    context = drafter.calls[0][1]
    assert "<tender_documents>" in context[0] and "Recruitment portal" in context[0]
    assert "Acme Softech" in context[1] and "CA certificate FY24" in context[1]
    assert "HRMS for State PSU" in context[1]  # similar previous proposal used as reference
    assert drafter.calls[1][1] is context  # same cached prefix for every call
    assert r["compliance"][0]["requirement"] == "Average turnover of Rs 2.3 lakh"
    assert r["proposal"]["slides"][0]["title"] == "Our solution"
    docs = {d["requirement"]: d["status"] for d in r["documents"]["required"]}
    assert docs["Undertaking of non-blacklisting"] == "prepare"


def test_ai_failure_keeps_other_results(env):
    a = _run(env, FakeDrafter(fail_on="compliance"))
    assert a.status == "done"
    assert a.result["stages"]["compliance"]["status"] == "failed"
    assert a.result["proposal"]["executive_summary"]
    assert any("simulated failure" in w for w in a.result["warnings"])


def test_exports_open_cleanly(env):
    a = _run(env, FakeDrafter())
    cfg = env[0]
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    doc = Document(io.BytesIO(to_docx(a.result, a.title, cfg.company)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Executive summary" in text and "Acme Softech" in text
    prs = Presentation(io.BytesIO(to_pptx(a.result, a.title, cfg.company)))
    assert len(prs.slides) == 2  # cover + 1 drafted slide
    wb = load_workbook(io.BytesIO(to_xlsx(a.result)))
    assert wb["Compliance matrix"]["E2"].value == "Complied"
    assert wb["BoQ"]["B2"].value == "Recruitment portal"
    # exports also work without any AI output
    plain = _run(env, None)
    assert to_docx(plain.result, plain.title, cfg.company) and to_pptx(plain.result, plain.title, cfg.company)


def test_boq_and_document_matching_helpers():
    rows = [["Sl. No", "Particulars", "Qty", "Unit"], ["1", "Annual maintenance", "12", "Months"], ["2", "Hosting", "1", "Lot"]]
    items = find_boq(rows, "upload")
    assert [i["title"] for i in items] == ["Annual maintenance", "Hosting"] and items[0]["quantity"] == 12
    assert find_boq([["random", "table"], ["no", "header"]]) == []
    assert boq_from_listing("Item A,Item B, with comma", None)[1]["title"] == "Item B, with comma"
    assert categorise_requirement("Company PAN card") == "pan"
    assert categorise_requirement("Japan office address") == "other"
    assert match_documents(["GST registration"], [])[0]["status"] == "missing"


def test_added_documents_are_read_and_unreadable_links_listed(env):
    cfg, Session, aid = env
    sow = cfg.data_dir / "sow.txt"
    sow.write_text("Scope of Work: deploy 11 resources on-site for 24 months at DESIDOC.", encoding="utf-8")
    with Session() as s:
        a = s.get(TenderAnalysis, aid)
        a.extra_files = [{"name": "SOW.txt", "path": str(sow)}, {"name": "gone.pdf", "path": str(cfg.data_dir / "missing.pdf")}]
        a.bid.attachments.append(BidAttachment(url="https://bidplus.gem.gov.in/resources/x/sow.pdf", label="Scope of Work",
                                               kind="sow", status="blocked_by_robots"))  # fmt: skip
        s.commit()
    drafter = FakeDrafter()
    r = _run(env, drafter).result
    corpus = drafter.calls[0][1][0]
    assert '<document name="Added by you: SOW.txt">' in corpus and "24 months at DESIDOC" in corpus
    assert r["unread_documents"][0]["label"] == "Scope of Work"
    assert any("gone.pdf" in w for w in r["warnings"])
    assert r["scope"]["facts"][0]["basis"] == "stated"


def test_pdf_text_carries_page_markers(tmp_path):
    from tender_tracker.pdf_parser import PdfContent

    content = PdfContent("a\nb", [], [], ["first page", "second (cid:3)page"])
    assert content.paged_text() == "[Page 1]\nfirst page\n[Page 2]\nsecond page"


def test_boq_fallback_uses_spec_and_consignee_tables():
    from tender_tracker.boq import boq_from_bid
    from tender_tracker.pdf_parser import BidDocumentInfo, apply_tables

    info = BidDocumentInfo()
    apply_tables(info, [
        [["/ Specification", "/ Values"], ["Type of AI solution", "Chatbots"], ["Input Data Type", "Audio , Image , Structured , Semi-Structured"],
         ["Model", "As per Service Provider's Solutions"]],
        [["..../S.N o.", "// /Consignee Reporting/Officer", "/Address", "Quantity", "/Additional Requirement"],
         ["1", "*****", "*****North Delhi", "Project / Lumpsum Based", "Number of Resources : 11"]],
    ])  # fmt: skip
    assert info.consignees == [{"address": "North Delhi", "quantity": "Project / Lumpsum Based", "additional": "Number of Resources : 11"}]
    sections = [[[r.group, r.name, r.value] for r in s] for s in info.spec_sections]
    [item] = boq_from_bid("AI as a Service - Development; Audio, Image, Structured, S..", 1, sections, info.consignees)
    assert item["title"].endswith("(name shortened by GeM)") and ".." not in item["title"]
    assert "Input Data Type: Audio , Image , Structured , Semi-Structured" in item["description"]
    assert "As per Service Provider" not in item["description"]  # generic values dropped
    assert "Number of Resources : 11" in item["description"]
    assert item["quantity_text"] == "Project / Lumpsum Based" and item["quantity"] is None
